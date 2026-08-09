"""Chargement du contenu textuel des documents locaux."""

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".txt", ".docx", ".pdf"}

IGNORED_KEYWORDS = {
    "archive",
    "archives",
    "ancien",
    "anciens",
    "old",
    "backup",
    "sauvegarde",
}


@dataclass
class LoadedDocument:
    """Document extrait avec ses métadonnées principales."""

    path: Path
    text: str
    category: str
    subcategory: str
    extension: str


def read_txt(path: Path) -> str:
    """Lit un fichier texte en essayant plusieurs encodages."""

    encodings = (
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin-1",
    )

    for encoding in encodings:
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue

    return path.read_text(
        encoding="utf-8",
        errors="replace",
    )


def read_docx(path: Path) -> str:
    """Extrait les paragraphes et tableaux d'un fichier DOCX."""

    document = Document(path)

    parts: list[str] = []

    # Paragraphes
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tableaux
    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if values:
                parts.append(" | ".join(values))

    return "\n".join(parts)


def read_pdf(path: Path) -> str:
    """Extrait le texte d'un PDF page par page."""

    reader = PdfReader(path)

    pages: list[str] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        text = page.extract_text()

        if text and text.strip():
            pages.append(
                f"[Page {page_number}]\n"
                f"{text.strip()}"
            )

    return "\n\n".join(pages)


def should_ignore_path(
    path: Path,
    documents_root: Path,
) -> bool:
    """
    Détermine si un fichier appartient à un dossier
    qui ne doit pas être indexé.
    """

    relative_parts = path.relative_to(
        documents_root
    ).parts[:-1]

    for part in relative_parts:
        normalized = part.strip().lower()

        if any(
            keyword in normalized
            for keyword in IGNORED_KEYWORDS
        ):
            return True

    return False


def extract_classification(
    path: Path,
    documents_root: Path,
) -> tuple[str, str]:
    """
    Déduit automatiquement la catégorie et la
    sous-catégorie depuis l'arborescence.
    """

    relative_path = path.relative_to(
        documents_root
    )

    parts = relative_path.parts

    # Exemple :
    #
    # documents/
    # └── FTTH/
    #     └── CGNAT/
    #         └── fichier.pdf
    #
    # category = FTTH
    # subcategory = CGNAT

    if len(parts) >= 2:
        category = parts[0]
    else:
        category = "Non classé"

    if len(parts) >= 3:
        subcategory = parts[1]
    else:
        subcategory = "Racine"

    return category, subcategory


def load_document(
    path: Path,
    documents_root: Path,
) -> LoadedDocument:
    """
    Charge le texte d'un document compatible et
    lui associe ses métadonnées.
    """

    extension = path.suffix.lower()

    if extension == ".txt":
        text = read_txt(path)

    elif extension == ".docx":
        text = read_docx(path)

    elif extension == ".pdf":
        text = read_pdf(path)

    else:
        raise ValueError(
            f"Format non pris en charge : {extension}"
        )

    category, subcategory = extract_classification(
        path=path,
        documents_root=documents_root,
    )

    return LoadedDocument(
        path=path,
        text=text,
        category=category,
        subcategory=subcategory,
        extension=extension,
    )


def find_documents(
    documents_root: Path,
) -> list[Path]:
    """
    Recherche récursivement tous les documents
    compatibles à indexer.
    """

    documents: list[Path] = []

    for path in documents_root.rglob("*"):

        # Ignorer les dossiers
        if not path.is_file():
            continue

        # Ignorer les fichiers temporaires Word
        # Exemple : ~$procedure.docx
        if path.name.startswith("~$"):
            continue

        # Format non encore pris en charge
        if (
            path.suffix.lower()
            not in SUPPORTED_EXTENSIONS
        ):
            continue

        # Archives, anciens documents, backups...
        if should_ignore_path(
            path=path,
            documents_root=documents_root,
        ):
            continue

        documents.append(path)

    return sorted(documents)